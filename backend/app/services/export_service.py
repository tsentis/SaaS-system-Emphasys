"""Render a list of projects to CSV, XLSX, PDF, or Word bytes."""

import csv
import io
from collections.abc import Sequence

from app.models.project import Project

COLUMNS = ["Acronym", "Title", "Status", "Total budget", "Start", "End", "Confidence"]


def _row(p: Project) -> list[str]:
    return [
        p.acronym or "",
        p.title or "",
        p.status or "",
        f"{p.total_budget:.2f}" if p.total_budget is not None else "",
        p.start_date.isoformat() if p.start_date else "",
        p.end_date.isoformat() if p.end_date else "",
        f"{float(p.extraction_confidence):.2f}" if p.extraction_confidence is not None else "",
    ]


def to_csv(projects: Sequence[Project]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(COLUMNS)
    for p in projects:
        writer.writerow(_row(p))
    return buf.getvalue().encode("utf-8")


def to_xlsx(projects: Sequence[Project]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Projects"
    ws.append(COLUMNS)
    for p in projects:
        ws.append(_row(p))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_docx(projects: Sequence[Project]) -> bytes:
    from docx import Document as Docx

    doc = Docx()
    doc.add_heading("European projects", level=1)
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Light Grid Accent 1"
    for i, col in enumerate(COLUMNS):
        table.rows[0].cells[i].text = col
    for p in projects:
        cells = table.add_row().cells
        for i, val in enumerate(_row(p)):
            cells[i].text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(projects: Sequence[Project]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4))
    data = [COLUMNS] + [_row(p) for p in projects]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
            ]
        )
    )
    doc.build([table])
    return buf.getvalue()
